"""Unit tests for the compressor application."""
import io
import json
import os
import shutil
import tempfile
import unittest
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.files.base import ContentFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from PIL import Image

from .models import CompressedFile, CompressionJob, UploadedFile, detect_category
from .services.converter import (
    ConversionUnavailable,
    conversion_available,
    convert_document,
    normalize_document_format,
    normalize_image_format,
)
from .services.dispatcher import process_job, resolve_document_extension, resolve_image_extension
from .services.image_compressor import compress_image
from .services.office_compressor import compress_office
from .services.pdf_compressor import compress_pdf
from .services.size_converter import (
    BYTES_PER_KB,
    BYTES_PER_MB,
    bytes_to_human,
    parse_size_string,
    parse_target_size,
)

MEDIA_ROOT = tempfile.mkdtemp(prefix="sc-test-media-")


def make_random_image(path, size=(1400, 1000), fmt="PNG"):
    width, height = size
    data = os.urandom(width * height * 3)
    image = Image.frombytes("RGB", (width, height), data)
    image.save(path, format=fmt)
    return path


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class SizeConverterTests(TestCase):
    def test_bytes_to_human(self):
        self.assertEqual(bytes_to_human(512), "512 B")
        self.assertEqual(bytes_to_human(BYTES_PER_KB), "1.00 KB")
        self.assertEqual(bytes_to_human(BYTES_PER_MB), "1.00 MB")

    def test_parse_size_string(self):
        self.assertEqual(parse_size_string("10 KB"), 10 * BYTES_PER_KB)
        self.assertEqual(parse_size_string("1.5 MB"), int(1.5 * BYTES_PER_MB))
        self.assertEqual(parse_size_string("2gb"), 2 * BYTES_PER_MB * 1024)
        self.assertIsNone(parse_size_string("nonsense"))

    def test_parse_target_size(self):
        self.assertEqual(parse_target_size(100, "KB"), 100 * BYTES_PER_KB)
        self.assertEqual(parse_target_size(2, "MB"), 2 * BYTES_PER_MB)
        self.assertIsNone(parse_target_size(0, "KB"))
        self.assertIsNone(parse_target_size("", "KB"))


class CategoryTests(TestCase):
    def test_detect_category(self):
        self.assertEqual(detect_category("photo.jpg"), UploadedFile.CATEGORY_IMAGE)
        self.assertEqual(detect_category("report.pdf"), UploadedFile.CATEGORY_DOCUMENT)
        self.assertEqual(detect_category("data.bin"), UploadedFile.CATEGORY_OTHER)


class ImageCompressorTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.source = os.path.join(self.tmp, "source.png")
        make_random_image(self.source)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_best_effort_reduces_size(self):
        output = os.path.join(self.tmp, "out.jpg")
        result = compress_image(self.source, output, level="medium")
        self.assertTrue(os.path.exists(output))
        self.assertLess(result["compressed_size"], result["original_size"])
        self.assertGreater(result["savings_percent"], 0)

    def test_target_size_is_met(self):
        output = os.path.join(self.tmp, "target.jpg")
        target = 60 * BYTES_PER_KB
        result = compress_image(self.source, output, target_bytes=target, level="medium")
        self.assertTrue(os.path.exists(output))
        self.assertLessEqual(result["compressed_size"], target)
        self.assertTrue(result["achieved_target"])

    def test_small_file_is_not_enlarged(self):
        small = os.path.join(self.tmp, "small.jpg")
        Image.new("RGB", (40, 40), (120, 120, 200)).save(small, "JPEG", quality=60)
        original_size = os.path.getsize(small)
        output = os.path.join(self.tmp, "small_out.jpg")
        result = compress_image(small, output, target_bytes=500 * BYTES_PER_KB)
        self.assertLessEqual(result["compressed_size"], original_size)

    def test_convert_to_webp(self):
        output = os.path.join(self.tmp, "out.webp")
        result = compress_image(self.source, output, convert_to_webp=True, level="medium")
        self.assertEqual(result["format"], "WEBP")
        self.assertTrue(os.path.exists(output))


class PDFCompressorTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        import pymupdf

        image_path = os.path.join(self.tmp, "img.jpg")
        make_random_image(image_path, fmt="JPEG")
        doc = pymupdf.open()
        page = doc.new_page(width=595, height=842)
        page.insert_image(pymupdf.Rect(40, 40, 555, 500), filename=image_path)
        page.insert_text((50, 560), "SmartCompress Pro PDF compression test.")
        self.pdf_path = os.path.join(self.tmp, "doc.pdf")
        doc.save(self.pdf_path)
        doc.close()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_pdf_compression(self):
        output = os.path.join(self.tmp, "out.pdf")
        result = compress_pdf(self.pdf_path, output, level="high")
        self.assertTrue(os.path.exists(output))
        self.assertGreater(result["original_size"], 0)
        self.assertLessEqual(result["compressed_size"], result["original_size"])


class OfficeCompressorTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        from docx import Document
        from docx.shared import Inches

        image_path = os.path.join(self.tmp, "img.jpg")
        make_random_image(image_path, fmt="JPEG")
        document = Document()
        document.add_heading("SmartCompress Pro", level=1)
        document.add_paragraph("This paragraph and its formatting must be preserved.")
        document.add_picture(image_path, width=Inches(5))
        self.docx_path = os.path.join(self.tmp, "report.docx")
        document.save(self.docx_path)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_docx_compression_preserves_package(self):
        output = os.path.join(self.tmp, "out.docx")
        result = compress_office(self.docx_path, output, level="high")
        self.assertTrue(os.path.exists(output))
        self.assertTrue(zipfile.is_zipfile(output))
        with zipfile.ZipFile(output) as archive:
            names = archive.namelist()
            self.assertIn("word/document.xml", names)
            self.assertIn("word/media/image1.jpg", names)
        self.assertLessEqual(result["compressed_size"], result["original_size"])


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class FormatConversionTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_normalize_helpers(self):
        self.assertEqual(normalize_image_format("JPEG"), "jpg")
        self.assertEqual(normalize_image_format("auto"), None)
        self.assertEqual(normalize_document_format(".PDF"), "pdf")
        self.assertEqual(normalize_document_format("keep"), None)

    def test_resolve_extensions(self):
        self.assertEqual(resolve_image_extension(".png", {"output_format": "webp"}), ".webp")
        self.assertEqual(resolve_image_extension(".png", {"output_format": "auto", "convert_to_webp": True}), ".webp")
        self.assertEqual(resolve_image_extension(".png", {"output_format": "auto"}), ".png")
        self.assertEqual(resolve_document_extension(".docx", {"output_format": "pdf"}), ".pdf")
        self.assertEqual(resolve_document_extension(".docx", {"output_format": "auto"}), ".docx")

    def test_image_format_conversion(self):
        source = os.path.join(self.tmp, "conv_src.png")
        make_random_image(source, size=(600, 400), fmt="PNG")
        result = compress_image(source, os.path.join(self.tmp, "conv_out.jpg"), level="medium")
        self.assertEqual(result["format"], "JPEG")

    def test_png_alpha_to_jpg_flattens(self):
        source = os.path.join(self.tmp, "alpha.png")
        Image.new("RGBA", (200, 200), (255, 0, 0, 128)).save(source, "PNG")
        output = os.path.join(self.tmp, "alpha.jpg")
        compress_image(source, output, level="medium")
        with Image.open(output) as img:
            self.assertEqual(img.mode, "RGB")

    def test_process_job_converts_image_format(self):
        buffer = io.BytesIO()
        make_random_image(buffer, size=(500, 400), fmt="PNG")
        image_bytes = buffer.getvalue()
        upload = UploadedFile.objects.create(
            original_file=ContentFile(image_bytes, name="pic.png"),
            filename="pic.png",
            file_type="png",
            category=UploadedFile.CATEGORY_IMAGE,
            original_size=len(image_bytes),
        )
        job = CompressionJob.objects.create(
            uploaded_file=upload,
            compression_level=CompressionJob.LEVEL_MEDIUM,
            options={"output_format": "jpg"},
        )
        process_job(job)
        job.refresh_from_db()
        self.assertEqual(job.status, CompressionJob.STATUS_COMPLETED)
        self.assertTrue(job.result.compressed_file.name.endswith(".jpg"))
        with Image.open(job.result.compressed_file.path) as img:
            self.assertEqual(img.format, "JPEG")

    def test_process_job_keeps_format_when_auto(self):
        buffer = io.BytesIO()
        make_random_image(buffer, size=(300, 300), fmt="PNG")
        image_bytes = buffer.getvalue()
        upload = UploadedFile.objects.create(
            original_file=ContentFile(image_bytes, name="keep.png"),
            filename="keep.png",
            file_type="png",
            category=UploadedFile.CATEGORY_IMAGE,
            original_size=len(image_bytes),
        )
        job = CompressionJob.objects.create(
            uploaded_file=upload,
            compression_level=CompressionJob.LEVEL_MEDIUM,
            options={"output_format": "auto"},
        )
        process_job(job)
        job.refresh_from_db()
        self.assertTrue(job.result.compressed_file.name.endswith(".png"))

    @unittest.skipIf(conversion_available(), "LibreOffice is installed")
    def test_document_conversion_requires_libreoffice(self):
        with self.assertRaises(ConversionUnavailable):
            convert_document("whatever.docx", "pdf")


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class ViewTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.image = io.BytesIO()
        Image.new("RGB", (60, 60), (10, 40, 90)).save(self.image, "PNG")
        self.image.seek(0)

    def test_home_page_loads(self):
        response = self.client.get(reverse("compressor:home"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Compress Smarter")

    def test_compressor_pages_load(self):
        self.assertEqual(self.client.get(reverse("compressor:image_compressor")).status_code, 200)
        self.assertEqual(self.client.get(reverse("compressor:document_compressor")).status_code, 200)

    def test_upload_and_compress_flow(self):
        upload = SimpleUploadedFile("test.png", self.image.read(), content_type="image/png")
        response = self.client.post(reverse("compressor:upload"), {"file": upload})
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertIn("id", payload)
        self.assertTrue(UploadedFile.objects.filter(pk=payload["id"]).exists())

        response = self.client.post(
            reverse("compressor:compress"),
            data='{"upload_ids": [%d], "target_value": 50, "target_unit": "KB", "level": "medium"}' % payload["id"],
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        jobs = response.json()["jobs"]
        self.assertEqual(len(jobs), 1)
        self.assertIn(jobs[0]["status"], {"completed", "warning"})
        self.assertTrue(CompressedFile.objects.filter(compression_job_id=jobs[0]["id"]).exists())

    def test_upload_rejects_unsupported_type(self):
        bad = SimpleUploadedFile("malicious.exe", b"MZ\x90\x00", content_type="application/octet-stream")
        response = self.client.post(reverse("compressor:upload"), {"file": bad})
        self.assertEqual(response.status_code, 400)

    def test_job_status_requires_access(self):
        response = self.client.get(reverse("compressor:job_status", args=[9999]))
        self.assertEqual(response.status_code, 404)

    def test_dashboard_requires_login(self):
        response = self.client.get(reverse("compressor:dashboard"))
        self.assertEqual(response.status_code, 302)

    def test_dashboard_renders_when_logged_in(self):
        from django.contrib.auth.models import User

        User.objects.create_user(username="dash", password="Sup3rSecret!")
        self.client.login(username="dash", password="Sup3rSecret!")
        response = self.client.get(reverse("compressor:dashboard"))
        self.assertEqual(response.status_code, 200)

    def test_results_page_renders(self):
        upload = SimpleUploadedFile("res.png", self.image.read(), content_type="image/png")
        upload_id = self.client.post(reverse("compressor:upload"), {"file": upload}).json()["id"]
        job = self.client.post(
            reverse("compressor:compress"),
            data='{"upload_ids": [%d], "level": "medium"}' % upload_id,
            content_type="application/json",
        ).json()["jobs"][0]
        response = self.client.get(reverse("compressor:results", args=[job["id"]]))
        self.assertEqual(response.status_code, 200)

    def test_download_and_zip_download(self):
        upload = SimpleUploadedFile("dl.png", self.image.read(), content_type="image/png")
        upload_id = self.client.post(reverse("compressor:upload"), {"file": upload}).json()["id"]
        jobs = self.client.post(
            reverse("compressor:compress"),
            data='{"upload_ids": [%d], "level": "medium"}' % upload_id,
            content_type="application/json",
        ).json()["jobs"]
        job_id = jobs[0]["id"]

        download = self.client.get(reverse("compressor:download", args=[job_id]))
        self.assertEqual(download.status_code, 200)
        self.assertIn("attachment", download["Content-Disposition"])

        archive = self.client.post(
            reverse("compressor:download_all"),
            data=json.dumps({"job_ids": [job_id]}),
            content_type="application/json",
        )
        self.assertEqual(archive.status_code, 200)
        self.assertEqual(archive["Content-Type"], "application/zip")


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class AuthFlowTests(TestCase):
    def test_register_creates_profile(self):
        response = self.client.post(
            reverse("accounts:register"),
            {
                "username": "newuser",
                "email": "new@example.com",
                "password1": "Sup3rSecret!",
                "password2": "Sup3rSecret!",
            },
        )
        self.assertEqual(response.status_code, 302)
        from django.contrib.auth.models import User

        user = User.objects.get(username="newuser")
        self.assertTrue(hasattr(user, "profile"))
